"""CLI: exporta advogados do polo ativo e polo passivo (uma linha por processo, uma aba)."""

from __future__ import annotations

import importlib
import sys
from collections.abc import Callable
from concurrent.futures import ProcessPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from types import ModuleType

_ROOT = Path(__file__).resolve().parents[1]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

import pandas as pd

from config.tribunais import TRIBUNAIS_MAP
from scrapers.consulta_cli_common import (
    escala_sleep_scraper_module,
    import_class,
    mp_worker_consulta_chunk,
    resolver_diretorio_saida_cli,
    sanitizar_dataframe_para_excel,
)
from scrapers.raspar_consulta_movimentacoes_core import parse_args_mov
from scrapers.raspar_consulta_planilhas_core import (
    _parse_aba_arg,
    _parse_coluna_arg,
    _split_indexed_chunks,
    ler_numeros_entrada,
    limpar_texto,
)


def _texto_advogado(p: dict) -> str:
    """Texto como no PJe (prioriza nome_completo da célula)."""
    nc = limpar_texto(p.get("nome_completo") or "")
    if nc:
        return nc
    partes: list[str] = []
    if p.get("nome"):
        partes.append(limpar_texto(p.get("nome")))
    if p.get("oab"):
        partes.append(f"OAB {limpar_texto(p.get('oab'))}")
    if p.get("cpf"):
        partes.append(f"CPF: {limpar_texto(p.get('cpf'))}")
    if p.get("cnpj"):
        partes.append(f"CNPJ: {limpar_texto(p.get('cnpj'))}")
    if p.get("tipo"):
        partes.append(f"({limpar_texto(p.get('tipo'))})")
    return " - ".join(partes)


def _es_linha_advogado(p: dict) -> bool:
    """Mesma heurística que PjePolosAdvMixin._linha_e_advogado (litigante vs adv)."""
    tipo = (p.get("tipo") or "").upper()
    texto = (p.get("nome_completo") or p.get("nome") or "").lower()
    if p.get("oab"):
        return True
    if "ADVOGADO" in tipo or "ADVOGADA" in tipo or "PROCURADOR" in tipo:
        return True
    if "advogado" in texto or "advogada" in texto or "procurador" in texto:
        return True
    return False


def _texto_parte_litigante(p: dict) -> str:
    """Parte que não é advogado (polo ativo/passivo)."""
    nc = limpar_texto(p.get("nome_completo") or "")
    if nc:
        return nc
    partes: list[str] = []
    if p.get("nome"):
        partes.append(limpar_texto(p.get("nome")))
    if p.get("cpf"):
        partes.append(f"CPF: {limpar_texto(p.get('cpf'))}")
    if p.get("cnpj"):
        partes.append(f"CNPJ: {limpar_texto(p.get('cnpj'))}")
    if p.get("tipo"):
        partes.append(f"({limpar_texto(p.get('tipo'))})")
    return " - ".join(partes)


def montar_planilha_adv_ativo_passivo(resultados: list[dict]) -> pd.DataFrame:
    """
    Uma linha por processo; colunas: Numero Processo | parte ativo | parte passivo |
    adv ativo | adv passivo.
    Partes = linhas do polo que não são advogados; advs = só advogados por polo.
    """
    colunas = [
        "Numero Processo",
        "parte ativo",
        "parte passivo",
        "adv ativo",
        "adv passivo",
    ]
    linhas: list[dict] = []
    for resultado in resultados:
        num = limpar_texto(resultado.get("numero_processo", ""))
        if resultado.get("status") != "sucesso" or not resultado.get("dados"):
            err = limpar_texto(resultado.get("erro") or "")
            linhas.append(
                {
                    "Numero Processo": num,
                    "parte ativo": err,
                    "parte passivo": "",
                    "adv ativo": "",
                    "adv passivo": "",
                }
            )
            continue
        d = resultado["dados"]
        polo_ativo = d.get("polo_ativo") or []
        polo_passivo = d.get("polo_passivo") or []

        partes_at_list = [_texto_parte_litigante(p) for p in polo_ativo if not _es_linha_advogado(p)]
        partes_at_list = [t for t in partes_at_list if t]
        partes_pa_list = [_texto_parte_litigante(p) for p in polo_passivo if not _es_linha_advogado(p)]
        partes_pa_list = [t for t in partes_pa_list if t]

        advs = d.get("advogados") or []
        ativos = [a for a in advs if (a.get("polo_participacao") or "").lower() == "ativo"]
        passivos = [a for a in advs if (a.get("polo_participacao") or "").lower() == "passivo"]
        linhas_at = [_texto_advogado(a) for a in ativos]
        linhas_at = [t for t in linhas_at if t]
        linhas_pa = [_texto_advogado(a) for a in passivos]
        linhas_pa = [t for t in linhas_pa if t]

        txt_parte_at = "\n".join(partes_at_list)
        txt_parte_pa = "\n".join(partes_pa_list)
        txt_adv_at = "\n".join(linhas_at)
        txt_adv_pa = "\n".join(linhas_pa)
        ext = limpar_texto(d.get("erro_extracao") or "")
        if ext and not txt_adv_at and not txt_adv_pa and not txt_parte_at and not txt_parte_pa:
            txt_parte_at = ext
        linhas.append(
            {
                "Numero Processo": num,
                "parte ativo": txt_parte_at,
                "parte passivo": txt_parte_pa,
                "adv ativo": txt_adv_at,
                "adv passivo": txt_adv_pa,
            }
        )
    return pd.DataFrame(linhas, columns=colunas)


def _payload_extras_pol(ch: ModuleType) -> dict:
    return {
        "tribunal_key": ch.TRIBUNAL_KEY,
        "sleep_module": ch.SLEEP_MODULE,
        "cli_class_path": ch.CLI_CLASS_PATH,
        "mov_class_path": ch.MOV_CLASS_PATH,
        "pol_class_path": ch.POL_CLASS_PATH,
    }


def main_polos(
    helpers_module: str,
    argv: list[str],
    *,
    progress_cb: Callable[[dict], None] | None = None,
) -> int:
    ch = importlib.import_module(helpers_module)
    pol_path = getattr(ch, "POL_CLASS_PATH", None)
    if not pol_path:
        print(
            f"Módulo {helpers_module} não define POL_CLASS_PATH (CLI polos não suportado para este tribunal).",
            file=sys.stderr,
        )
        return 1

    args = parse_args_mov(argv, getattr(ch, "DESC_POL", "Exporta polos e advogados para Excel."))
    numeros: list[str] = list(args.numeros or [])
    coluna = _parse_coluna_arg(args.coluna)
    aba = _parse_aba_arg(args.aba)

    if args.entrada:
        if not args.entrada.is_file():
            print(f"Arquivo não encontrado: {args.entrada}", file=sys.stderr)
            return 1
        try:
            numeros.extend(ler_numeros_entrada(args.entrada, aba, coluna))
        except Exception as e:
            print(f"Erro ao ler entrada: {e}", file=sys.stderr)
            return 1

    seen: set[str] = set()
    ordem: list[str] = []
    for n in numeros:
        if n not in seen:
            seen.add(n)
            ordem.append(n)
    numeros = ordem

    if not numeros:
        print("Use --entrada e/ou --numero.", file=sys.stderr)
        return 1

    if args.headless and args.sem_headless:
        print("Use só uma: --headless ou --sem-headless.", file=sys.stderr)
        return 1

    cfg = dict(TRIBUNAIS_MAP[ch.TRIBUNAL_KEY])
    if args.headless:
        cfg["headless"] = True
    elif args.sem_headless:
        cfg["headless"] = False

    saida_dir = resolver_diretorio_saida_cli(args.saida_dir, args.entrada)
    if not saida_dir.is_absolute():
        saida_dir = _ROOT / saida_dir
    saida_dir.mkdir(parents=True, exist_ok=True)
    slug = getattr(ch, "SLUG", ch.TRIBUNAL_KEY.replace(".", "_"))
    if args.entrada and args.entrada.is_file():
        path_out = saida_dir / f"{Path(args.entrada).stem}_polos_advogados.xlsx"
    else:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        path_out = saida_dir / f"polos_advogados_{slug}_{ts}.xlsx"

    if args.sleep_scale is not None:
        fator_sleep = max(0.15, min(1.0, float(args.sleep_scale)))
    elif args.rapido:
        fator_sleep = 0.55
    else:
        fator_sleep = 1.0

    workers = max(1, min(int(args.workers or 1), 8))
    if int(args.workers or 1) > 8:
        print("Aviso: --workers limitado a 8.", file=sys.stderr)

    resultados: list[dict] = []
    extras = _payload_extras_pol(ch)
    PolCls = import_class(ch.POL_CLASS_PATH)

    if workers == 1:
        with escala_sleep_scraper_module(ch.SLEEP_MODULE, fator_sleep):
            scraper = PolCls(cfg)
            try:
                scraper.setup_driver()
                for i, num in enumerate(numeros, 1):
                    print(f"[{i}/{len(numeros)}] {num} ...")
                    raw = scraper.raspar_processo(num)
                    res = ch.normalizar_resultado_scraper(num, raw)
                    resultados.append(res)
                    if progress_cb:
                        progress_cb(res)
                    print("  OK" if res["status"] == "sucesso" else f"  Erro: {res.get('erro')}")
            finally:
                scraper.close()
    else:
        indexed = list(enumerate(numeros))
        chunks = _split_indexed_chunks(indexed, workers)
        n_chunks = len(chunks)
        print(f"Paralelo: {n_chunks} workers, {len(numeros)} processos (partes e advogados por polo).")
        root_str = str(_ROOT)
        payloads = [
            {
                "root": root_str,
                "chunk": [[i, n] for i, n in chunk],
                "cfg": cfg,
                "fator_sleep": fator_sleep,
                "worker_id": wi + 1,
                "somente_polos_adv": True,
                **extras,
            }
            for wi, chunk in enumerate(chunks)
        ]
        slots: list[dict | None] = [None] * len(numeros)
        with ProcessPoolExecutor(max_workers=n_chunks) as pool:
            futures = [pool.submit(mp_worker_consulta_chunk, p) for p in payloads]
            for fut in as_completed(futures):
                for idx, res in fut.result():
                    slots[idx] = res
                    if progress_cb:
                        progress_cb(res)
        resultados = [slots[i] for i in range(len(numeros))]

    df = sanitizar_dataframe_para_excel(montar_planilha_adv_ativo_passivo(resultados))

    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(str(path_out), engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Polos e advogados")
        ws = writer.sheets["Polos e advogados"]
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col_idx, col_name in enumerate(df.columns, start=1):
            letter = get_column_letter(col_idx)
            if col_name == "Numero Processo":
                ws.column_dimensions[letter].width = 30
            else:
                ws.column_dimensions[letter].width = 85

    print(f"Planilha gerada: {path_out}")
    return 0

"""Núcleo do CLI de planilhas (raspada + tratada) parametrizado por módulo `cli_helpers` do tribunal."""

from __future__ import annotations

import argparse
import importlib
import re
import sys
from collections.abc import Callable
from concurrent.futures import ProcessPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from types import ModuleType

_ROOT = Path(__file__).resolve().parents[1]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

import pandas as pd

from app.services.transformador_dados import TransformadorDados
from config.tribunais import TRIBUNAIS_MAP
from scrapers.consulta_cli_common import (
    escala_sleep_scraper_module,
    import_class,
    mp_worker_consulta_chunk,
    resolver_diretorio_saida_cli,
    sanitizar_dataframe_para_excel,
)

COLUNAS_TRATADAS_ORDEM = [
    "pasta",
    "numeroProcessoAnterior",
    "cnj",
    "tipoPartePoloAtivo",
    "partePoloAtivo",
    "Advogado Parte Contraria",
    "tipoPartePoloPassivo",
    "partePoloPassivo",
    "cliente",
    "tipoDeRito",
    "dataDistribuicao",
    "numeroUnidade",
    "unidade",
    "especialidade",
    "comarca",
    "estado",
    "orgao",
    "natureza",
    "materia",
    "dataInstancia",
    "tipoInstancia",
    "sistemaExterno",
    "processoEletronico",
    "processoEstrategico",
    "valorCausa",
    "valorFinalCausa",
    "tipoAcao",
    "tipoObjeto",
    "dataFase",
    "fase",
    "dataStatus",
    "status",
    "grupoProcesso",
    "prioridadeDe",
    "data_resultado",
    "tipo_resultado",
    "descricao_resultado",
    "dataEvento",
    "tipoEvento",
    "descricaoEvento",
    "complementoEvento",
    "observacaoEvento",
    "solicitanteEvento",
    "responsavelEvento",
    "grupoTrabalho",
    "corresponsavel",
    "dataNotificacao",
    "dataNotificacaoAdicional",
    "probabilidadePerda",
    "dataValorProvisionado",
    "valorProvisionado",
    "dataAndamento",
    "tipoAndamento",
    "descricaoAndamento",
    "complementoAndamento",
    "solicitanteAndamento",
    "responsavelAndamento",
    "corresponsavelAndamento",
    "descricaoObjeto",
    "escritorioCredenciado",
    "dataContratacao",
    "observacaoDoProcesso",
    "parecerDoProcesso",
]


def limpar_texto(texto) -> str:
    if not texto:
        return ""
    texto = re.sub(r"\s+", " ", str(texto))
    return texto.strip()


def montar_linhas_planilha_raspada(resultados: list[dict]) -> list[dict]:
    dados_export = []
    for resultado in resultados:
        if resultado.get("status") == "sucesso" and resultado.get("dados"):
            dados = resultado["dados"]
            dados_processo = dados.get("dados_processo", {})
            polo_ativo = dados.get("polo_ativo", [])
            polo_passivo = dados.get("polo_passivo", [])
            movimentacoes = dados.get("movimentacoes", [])
            texto_movimentacoes = "\n".join(
                [
                    mov.get("movimento", "") or mov.get("descricao", "")
                    for mov in movimentacoes
                    if mov.get("movimento") or mov.get("descricao")
                ]
            )
            nomes_polo_ativo = [
                TransformadorDados.normalizar_nome_participante(
                    p.get("nome", "") or p.get("nome_completo", "")
                )
                for p in polo_ativo
            ]
            nomes_polo_passivo = [
                TransformadorDados.normalizar_nome_participante(
                    p.get("nome", "") or p.get("nome_completo", "")
                )
                for p in polo_passivo
            ]
            dados_export.append(
                {
                    "Numero Processo": limpar_texto(resultado.get("numero_processo", "")),
                    "Tribunal": limpar_texto(resultado.get("tribunal", "")),
                    "Data Distribuicao": limpar_texto(dados_processo.get("data_distribuicao", "")),
                    "Classe Judicial": limpar_texto(dados_processo.get("classe_judicial", "")),
                    "Assunto": limpar_texto(dados_processo.get("assunto", "")),
                    "Jurisdicao": limpar_texto(dados_processo.get("jurisdicao", "")),
                    "Orgao Julgador": limpar_texto(dados_processo.get("orgao_julgador", "")),
                    "Polo Ativo": "; ".join(n for n in nomes_polo_ativo if n),
                    "Polo Passivo": "; ".join(n for n in nomes_polo_passivo if n),
                    "Movimentacoes": limpar_texto(texto_movimentacoes),
                    "Total Movimentacoes": len(movimentacoes),
                    "Total Documentos": len(dados.get("documentos", [])),
                    "Erro": "",
                }
            )
        elif resultado.get("status") == "erro":
            dados_export.append(
                {
                    "Numero Processo": limpar_texto(resultado.get("numero_processo", "")),
                    "Tribunal": limpar_texto(resultado.get("tribunal", "")),
                    "Data Distribuicao": "",
                    "Classe Judicial": "",
                    "Assunto": "",
                    "Jurisdicao": "",
                    "Orgao Julgador": "",
                    "Polo Ativo": "",
                    "Polo Passivo": "",
                    "Movimentacoes": "",
                    "Total Movimentacoes": 0,
                    "Total Documentos": 0,
                    "Erro": limpar_texto(resultado.get("erro") or ""),
                }
            )
    return dados_export


def aplicar_formato_aba_raspada(worksheet) -> None:
    from openpyxl.styles import Alignment

    letters = "ABCDEFGHIJKLM"
    widths = [20, 15, 15, 40, 50, 30, 40, 40, 40, 80, 15, 15, 55]
    for letter, w in zip(letters, widths):
        worksheet.column_dimensions[letter].width = w
    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def aplicar_formato_aba_tratada(worksheet, df: pd.DataFrame) -> None:
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    colunas_sem_quebra = [
        "cnj",
        "numeroProcessoAnterior",
        "tipoPartePoloAtivo",
        "tipoPartePoloPassivo",
        "tipoDeRito",
        "dataDistribuicao",
        "numeroUnidade",
        "especialidade",
        "comarca",
        "estado",
        "natureza",
        "materia",
        "tipoInstancia",
        "processoEletronico",
        "processoEstrategico",
        "tipoAcao",
        "dataStatus",
        "status",
        "tipoEvento",
    ]
    for idx, col in enumerate(df.columns, 1):
        col_letter = get_column_letter(idx)
        if col == "cnj":
            worksheet.column_dimensions[col_letter].width = 25
        elif col in [
            "partePoloAtivo",
            "partePoloPassivo",
            "Advogado Parte Contraria",
            "descricaoEvento",
            "descricaoAndamento",
        ]:
            worksheet.column_dimensions[col_letter].width = 40
        else:
            worksheet.column_dimensions[col_letter].width = 20
    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
        for idx, cell in enumerate(row, 1):
            col_name = df.columns[idx - 1] if idx <= len(df.columns) else ""
            wrap_text = col_name not in colunas_sem_quebra
            cell.alignment = Alignment(wrap_text=wrap_text, vertical="top")


def ler_numeros_txt(path: Path) -> list[str]:
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    out = []
    for line in lines:
        s = line.strip()
        if s and not s.startswith("#"):
            out.append(s)
    return out


def ler_numeros_excel(path: Path, sheet: str | int | None, coluna: str | int) -> list[str]:
    df = pd.read_excel(path, sheet_name=sheet if sheet is not None else 0, header=0, dtype=str)
    if isinstance(coluna, int):
        if coluna < 0 or coluna >= len(df.columns):
            raise ValueError(f"Índice de coluna inválido: {coluna} (planilha tem {len(df.columns)} colunas).")
        serie = df.iloc[:, coluna]
    else:
        if coluna not in df.columns:
            raise ValueError(
                f"Coluna '{coluna}' não encontrada. Colunas: {list(df.columns)}"
            )
        serie = df[coluna]
    out = []
    for v in serie:
        if v is None or (isinstance(v, float) and pd.isna(v)):
            continue
        s = limpar_texto(v)
        if s:
            out.append(s)
    return out


def ler_numeros_entrada(path: Path, sheet: str | int | None, coluna: str | int | None) -> list[str]:
    suf = path.suffix.lower()
    if suf in (".xlsx", ".xls"):
        if coluna is None:
            coluna = 0
        return ler_numeros_excel(path, sheet, coluna)
    if suf == ".csv":
        df = pd.read_csv(path, dtype=str)
        if coluna is None:
            coluna = 0
        if isinstance(coluna, int):
            serie = df.iloc[:, coluna]
        else:
            serie = df[coluna]
        return [limpar_texto(v) for v in serie if limpar_texto(str(v) if v is not None else "")]
    return ler_numeros_txt(path)


def parse_args(argv: list[str], description: str) -> argparse.Namespace:
    p = argparse.ArgumentParser(description=description)
    p.add_argument("--entrada", type=Path, help="Arquivo .txt (um processo por linha), .csv ou .xlsx")
    p.add_argument(
        "--aba",
        default=None,
        help="Nome ou índice da aba Excel (padrão: primeira aba)",
    )
    p.add_argument(
        "--coluna",
        default=None,
        help="Excel/CSV: nome da coluna ou índice 0-based (padrão: primeira coluna)",
    )
    p.add_argument("--numero", action="append", dest="numeros", default=[], help="Número CNJ (repetir para vários)")
    p.add_argument(
        "--saida-dir",
        type=Path,
        default=None,
        help="Pasta dos .xlsx gerados. Omitido com --entrada: mesma pasta do ficheiro de entrada.",
    )
    p.add_argument("--headless", action="store_true", help="Força headless")
    p.add_argument("--sem-headless", action="store_true", help="Força com janela do Chrome")
    p.add_argument(
        "--rapido",
        action="store_true",
        help="Reduz pausas fixas do scraper (~55 por cento do tempo; consulta lenta pode dar timeout)",
    )
    p.add_argument(
        "--sleep-scale",
        type=float,
        default=None,
        metavar="F",
        help="Multiplica cada time.sleep do scraper (0.15–1.0). Ex.: 0.5 = metade. Vale sobre --rapido.",
    )
    p.add_argument(
        "--workers",
        type=int,
        default=1,
        metavar="N",
        help="Quantidade de processos Chrome em paralelo (1–8). Padrão 1 = sequencial. Mais RAM.",
    )
    return p.parse_args(argv)


def _parse_coluna_arg(val: str | None) -> str | int | None:
    if val is None:
        return None
    if val.isdigit():
        return int(val)
    return val


def _parse_aba_arg(val: str | None) -> str | int | None:
    if val is None:
        return None
    if val.isdigit():
        return int(val)
    return val


def _split_indexed_chunks(indexed: list[tuple[int, str]], num_workers: int) -> list[list[tuple[int, str]]]:
    n = len(indexed)
    if n == 0:
        return []
    if num_workers <= 1:
        return [indexed]
    nw = min(num_workers, n)
    q, r = divmod(n, nw)
    chunks: list[list[tuple[int, str]]] = []
    pos = 0
    for w in range(nw):
        sz = q + (1 if w < r else 0)
        chunks.append(indexed[pos : pos + sz])
        pos += sz
    return [c for c in chunks if c]


def _payload_extras(ch: ModuleType) -> dict:
    return {
        "tribunal_key": ch.TRIBUNAL_KEY,
        "sleep_module": ch.SLEEP_MODULE,
        "cli_class_path": ch.CLI_CLASS_PATH,
        "mov_class_path": ch.MOV_CLASS_PATH,
    }


def main_planilhas(
    helpers_module: str,
    argv: list[str],
    *,
    progress_cb: Callable[[dict], None] | None = None,
) -> int:
    ch = importlib.import_module(helpers_module)
    args = parse_args(argv, getattr(ch, "DESC_PLANILHAS", "Raspa consulta pública e gera planilhas raspada + tratada."))
    numeros: list[str] = list(args.numeros or [])
    coluna = _parse_coluna_arg(args.coluna)
    aba = _parse_aba_arg(args.aba)

    if args.entrada:
        if not args.entrada.is_file():
            print(f"Arquivo não encontrado: {args.entrada}", file=sys.stderr)
            return 1
        try:
            numeros.extend(ler_numeros_entrada(args.entrada, aba, coluna))
        except Exception as e:
            print(f"Erro ao ler entrada: {e}", file=sys.stderr)
            return 1

    seen: set[str] = set()
    ordem: list[str] = []
    for n in numeros:
        if n not in seen:
            seen.add(n)
            ordem.append(n)
    numeros = ordem

    if not numeros:
        print("Use --entrada (txt/csv/xlsx) e/ou --numero.", file=sys.stderr)
        return 1

    if args.headless and args.sem_headless:
        print("Use só uma: --headless ou --sem-headless.", file=sys.stderr)
        return 1

    cfg = dict(TRIBUNAIS_MAP[ch.TRIBUNAL_KEY])
    if args.headless:
        cfg["headless"] = True
    elif args.sem_headless:
        cfg["headless"] = False

    saida_dir = resolver_diretorio_saida_cli(args.saida_dir, args.entrada)
    if not saida_dir.is_absolute():
        saida_dir = _ROOT / saida_dir
    saida_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    slug = getattr(ch, "SLUG", ch.TRIBUNAL_KEY.replace(".", "_"))
    path_raspada = saida_dir / f"resultados_raspados_{slug}_{ts}.xlsx"
    path_tratada = saida_dir / f"resultados_tratados_{slug}_{ts}.xlsx"

    if args.sleep_scale is not None:
        fator_sleep = max(0.15, min(1.0, float(args.sleep_scale)))
    elif args.rapido:
        fator_sleep = 0.55
    else:
        fator_sleep = 1.0

    if fator_sleep < 1.0:
        print(
            f"Pausas do scraper em {fator_sleep:.0%} do tempo original (--rapido / --sleep-scale). "
            "Se aparecer muito timeout, rode sem essas opções ou use --sleep-scale mais alto (ex.: 0.75)."
        )

    workers = max(1, int(getattr(args, "workers", 1) or 1))
    if workers > 8:
        print("Aviso: --workers limitado a 8 neste CLI.", file=sys.stderr)
        workers = 8

    resultados: list[dict] = []
    extras = _payload_extras(ch)
    ScraperCls = import_class(ch.CLI_CLASS_PATH)

    if workers == 1:
        with escala_sleep_scraper_module(ch.SLEEP_MODULE, fator_sleep):
            scraper = ScraperCls(cfg)
            try:
                scraper.setup_driver()
                for i, num in enumerate(numeros, 1):
                    print(f"[{i}/{len(numeros)}] Raspando {num} ...")
                    raw = scraper.raspar_processo(num)
                    res = ch.normalizar_resultado_scraper(num, raw)
                    resultados.append(res)
                    if progress_cb:
                        progress_cb(res)
                    if res["status"] == "erro":
                        print(f"  Erro: {res.get('erro')}")
                    else:
                        print("  OK")
            finally:
                scraper.close()
    else:
        indexed = list(enumerate(numeros))
        chunks = _split_indexed_chunks(indexed, workers)
        n_chunks = len(chunks)
        print(
            f"Modo paralelo: {n_chunks} workers (Chrome), {len(numeros)} processos no total. "
            "Cada worker usa bastante RAM; se o sistema ficar instável, reduza --workers."
        )
        root_str = str(_ROOT)
        payloads = [
            {
                "root": root_str,
                "chunk": [[i, n] for i, n in ch],
                "cfg": cfg,
                "fator_sleep": fator_sleep,
                "worker_id": wi + 1,
                "somente_movimentacoes": False,
                **extras,
            }
            for wi, ch in enumerate(chunks)
        ]
        slots: list[dict | None] = [None] * len(numeros)
        with ProcessPoolExecutor(max_workers=n_chunks) as pool:
            futures = [pool.submit(mp_worker_consulta_chunk, p) for p in payloads]
            for fut in as_completed(futures):
                batch = fut.result()
                for idx, res in batch:
                    slots[idx] = res
                    if progress_cb:
                        progress_cb(res)
        resultados = [slots[i] for i in range(len(numeros))]

    df_raspada = sanitizar_dataframe_para_excel(pd.DataFrame(montar_linhas_planilha_raspada(resultados)))
    with pd.ExcelWriter(str(path_raspada), engine="openpyxl") as writer:
        df_raspada.to_excel(writer, index=False, sheet_name="Resultados")
        if not df_raspada.empty:
            aplicar_formato_aba_raspada(writer.sheets["Resultados"])

    dados_tratados = TransformadorDados.transformar_lote(resultados)
    df_tratada = pd.DataFrame(dados_tratados)
    if not df_tratada.empty:
        cols_ok = [c for c in COLUNAS_TRATADAS_ORDEM if c in df_tratada.columns]
        df_tratada = df_tratada[cols_ok]
    df_tratada = sanitizar_dataframe_para_excel(df_tratada)

    with pd.ExcelWriter(str(path_tratada), engine="openpyxl") as writer:
        df_tratada.to_excel(writer, index=False, sheet_name="Processos")
        if not df_tratada.empty:
            aplicar_formato_aba_tratada(writer.sheets["Processos"], df_tratada)

    print(f"Planilha raspada:  {path_raspada}")
    print(f"Planilha tratada: {path_tratada}")
    if df_raspada.empty:
        print("Aviso: planilha raspada sem linhas (nenhum sucesso).", file=sys.stderr)
    if df_tratada.empty:
        print("Aviso: planilha tratada vazia.", file=sys.stderr)
    return 0

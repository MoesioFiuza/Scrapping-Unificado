"""Núcleo do CLI de exportação só movimentações (uma linha por processo)."""

from __future__ import annotations

import argparse
import importlib
import sys
from collections.abc import Callable
from concurrent.futures import ProcessPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from types import ModuleType

_ROOT = Path(__file__).resolve().parents[1]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

import pandas as pd

from config.tribunais import TRIBUNAIS_MAP
from scrapers.consulta_cli_common import (
    escala_sleep_scraper_module,
    import_class,
    mp_worker_consulta_chunk,
    resolver_diretorio_saida_cli,
    sanitizar_dataframe_para_excel,
)
from scrapers.raspar_consulta_planilhas_core import (
    _parse_aba_arg,
    _parse_coluna_arg,
    _split_indexed_chunks,
    ler_numeros_entrada,
    limpar_texto,
)

MAX_PARES_MOVIMENTACAO_COLUNAS = 400


def montar_planilha_um_processo_por_linha(resultados: list[dict]) -> pd.DataFrame:
    max_n = 0
    for resultado in resultados:
        if resultado.get("status") == "sucesso" and resultado.get("dados"):
            movs = resultado["dados"].get("movimentacoes") or []
            max_n = max(max_n, len(movs))

    max_n = min(max_n, MAX_PARES_MOVIMENTACAO_COLUNAS)
    if max_n == MAX_PARES_MOVIMENTACAO_COLUNAS:
        real_max = 0
        for resultado in resultados:
            if resultado.get("status") == "sucesso" and resultado.get("dados"):
                real_max = max(real_max, len(resultado["dados"].get("movimentacoes") or []))
        if real_max > MAX_PARES_MOVIMENTACAO_COLUNAS:
            print(
                f"Aviso: há processos com mais de {MAX_PARES_MOVIMENTACAO_COLUNAS} movimentações; "
                f"colunas extras foram cortadas (máx. {MAX_PARES_MOVIMENTACAO_COLUNAS} pares).",
                file=sys.stderr,
            )

    colunas = ["Numero Processo", "Erro"]
    for i in range(1, max_n + 1):
        colunas.append(f"Data movimentação {i}")
        colunas.append(f"Movimentação {i}")

    linhas: list[dict] = []
    for resultado in resultados:
        row = {c: "" for c in colunas}
        num = limpar_texto(resultado.get("numero_processo", ""))
        row["Numero Processo"] = num

        if resultado.get("status") != "sucesso" or not resultado.get("dados"):
            row["Erro"] = limpar_texto(resultado.get("erro") or "")
            linhas.append(row)
            continue

        movs = resultado["dados"].get("movimentacoes") or []
        row["Erro"] = ""
        for j, m in enumerate(movs):
            if j >= max_n:
                break
            k = j + 1
            data = limpar_texto(m.get("data") or "")
            hora = limpar_texto(m.get("hora") or "")
            if data and hora:
                data_txt = f"{data} {hora}"
            else:
                data_txt = data or hora
            desc = limpar_texto(m.get("descricao") or "")
            mov_txt = limpar_texto(m.get("movimento") or "")
            texto = desc if desc else mov_txt
            row[f"Data movimentação {k}"] = data_txt
            row[f"Movimentação {k}"] = texto
        linhas.append(row)

    return pd.DataFrame(linhas, columns=colunas)


def montar_resumo_por_processo(resultados: list[dict]) -> list[dict]:
    out = []
    for resultado in resultados:
        num = limpar_texto(resultado.get("numero_processo", ""))
        if resultado.get("status") != "sucesso" or not resultado.get("dados"):
            out.append(
                {
                    "Numero Processo": num,
                    "Status": "erro",
                    "Total movimentacoes": 0,
                    "Erro": limpar_texto(resultado.get("erro") or ""),
                }
            )
        else:
            n = len(resultado["dados"].get("movimentacoes") or [])
            out.append(
                {
                    "Numero Processo": num,
                    "Status": "sucesso",
                    "Total movimentacoes": n,
                    "Erro": "",
                }
            )
    return out


def parse_args_mov(argv: list[str], description: str) -> argparse.Namespace:
    p = argparse.ArgumentParser(description=description)
    p.add_argument("--entrada", type=Path, help=".txt / .csv / .xlsx com números de processo")
    p.add_argument("--aba", default=None, help="Aba Excel (nome ou índice)")
    p.add_argument("--coluna", default=None, help="Coluna Excel/CSV (índice ou nome)")
    p.add_argument("--numero", action="append", dest="numeros", default=[])
    p.add_argument(
        "--saida-dir",
        type=Path,
        default=None,
        help="Pasta do .xlsx gerado. Omitido com --entrada: mesma pasta do ficheiro de entrada.",
    )
    p.add_argument("--headless", action="store_true")
    p.add_argument("--sem-headless", action="store_true")
    p.add_argument("--rapido", action="store_true")
    p.add_argument("--sleep-scale", type=float, default=None, metavar="F")
    p.add_argument("--workers", type=int, default=1, metavar="N")
    return p.parse_args(argv)


def _payload_extras(ch: ModuleType) -> dict:
    return {
        "tribunal_key": ch.TRIBUNAL_KEY,
        "sleep_module": ch.SLEEP_MODULE,
        "cli_class_path": ch.CLI_CLASS_PATH,
        "mov_class_path": ch.MOV_CLASS_PATH,
    }


def main_movimentacoes(
    helpers_module: str,
    argv: list[str],
    *,
    progress_cb: Callable[[dict], None] | None = None,
) -> int:
    ch = importlib.import_module(helpers_module)
    args = parse_args_mov(argv, getattr(ch, "DESC_MOV", "Exporta só movimentações para Excel."))
    numeros: list[str] = list(args.numeros or [])
    coluna = _parse_coluna_arg(args.coluna)
    aba = _parse_aba_arg(args.aba)

    if args.entrada:
        if not args.entrada.is_file():
            print(f"Arquivo não encontrado: {args.entrada}", file=sys.stderr)
            return 1
        try:
            numeros.extend(ler_numeros_entrada(args.entrada, aba, coluna))
        except Exception as e:
            print(f"Erro ao ler entrada: {e}", file=sys.stderr)
            return 1

    seen: set[str] = set()
    ordem: list[str] = []
    for n in numeros:
        if n not in seen:
            seen.add(n)
            ordem.append(n)
    numeros = ordem

    if not numeros:
        print("Use --entrada e/ou --numero.", file=sys.stderr)
        return 1

    if args.headless and args.sem_headless:
        print("Use só uma: --headless ou --sem-headless.", file=sys.stderr)
        return 1

    cfg = dict(TRIBUNAIS_MAP[ch.TRIBUNAL_KEY])
    if args.headless:
        cfg["headless"] = True
    elif args.sem_headless:
        cfg["headless"] = False

    saida_dir = resolver_diretorio_saida_cli(args.saida_dir, args.entrada)
    if not saida_dir.is_absolute():
        saida_dir = _ROOT / saida_dir
    saida_dir.mkdir(parents=True, exist_ok=True)
    slug = getattr(ch, "SLUG", ch.TRIBUNAL_KEY.replace(".", "_"))
    # Com ficheiro de entrada: nome da planilha (stem) + sufixo pedido; só --numero: padrão antigo com data.
    if args.entrada and args.entrada.is_file():
        path_out = saida_dir / f"{Path(args.entrada).stem}_movimentação.xlsx"
    else:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        path_out = saida_dir / f"movimentacoes_{slug}_{ts}.xlsx"

    if args.sleep_scale is not None:
        fator_sleep = max(0.15, min(1.0, float(args.sleep_scale)))
    elif args.rapido:
        fator_sleep = 0.55
    else:
        fator_sleep = 1.0

    workers = max(1, min(int(args.workers or 1), 8))
    if int(args.workers or 1) > 8:
        print("Aviso: --workers limitado a 8.", file=sys.stderr)

    resultados: list[dict] = []
    extras = _payload_extras(ch)
    MovCls = import_class(ch.MOV_CLASS_PATH)

    if workers == 1:
        with escala_sleep_scraper_module(ch.SLEEP_MODULE, fator_sleep):
            scraper = MovCls(cfg)
            try:
                scraper.setup_driver()
                for i, num in enumerate(numeros, 1):
                    print(f"[{i}/{len(numeros)}] {num} ...")
                    raw = scraper.raspar_processo(num)
                    res = ch.normalizar_resultado_scraper(num, raw)
                    resultados.append(res)
                    if progress_cb:
                        progress_cb(res)
                    print("  OK" if res["status"] == "sucesso" else f"  Erro: {res.get('erro')}")
            finally:
                scraper.close()
    else:
        indexed = list(enumerate(numeros))
        chunks = _split_indexed_chunks(indexed, workers)
        n_chunks = len(chunks)
        print(f"Paralelo: {n_chunks} workers, {len(numeros)} processos (só movimentações).")
        root_str = str(_ROOT)
        payloads = [
            {
                "root": root_str,
                "chunk": [[i, n] for i, n in ch],
                "cfg": cfg,
                "fator_sleep": fator_sleep,
                "worker_id": wi + 1,
                "somente_movimentacoes": True,
                **extras,
            }
            for wi, ch in enumerate(chunks)
        ]
        slots: list[dict | None] = [None] * len(numeros)
        with ProcessPoolExecutor(max_workers=n_chunks) as pool:
            futures = [pool.submit(mp_worker_consulta_chunk, p) for p in payloads]
            for fut in as_completed(futures):
                for idx, res in fut.result():
                    slots[idx] = res
                    if progress_cb:
                        progress_cb(res)
        resultados = [slots[i] for i in range(len(numeros))]

    df_mov = sanitizar_dataframe_para_excel(montar_planilha_um_processo_por_linha(resultados))
    df_resumo = sanitizar_dataframe_para_excel(pd.DataFrame(montar_resumo_por_processo(resultados)))

    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(str(path_out), engine="openpyxl") as writer:
        df_mov.to_excel(writer, index=False, sheet_name="Movimentacoes")
        df_resumo.to_excel(writer, index=False, sheet_name="Resumo")
        for name, df_sheet in (("Movimentacoes", df_mov), ("Resumo", df_resumo)):
            ws = writer.sheets[name]
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            for col_idx, col_name in enumerate(df_sheet.columns, start=1):
                letter = get_column_letter(col_idx)
                if name == "Movimentacoes":
                    if col_name == "Numero Processo":
                        ws.column_dimensions[letter].width = 28
                    elif col_name == "Erro":
                        ws.column_dimensions[letter].width = 36
                    elif str(col_name).startswith("Data movimentação"):
                        ws.column_dimensions[letter].width = 20
                    else:
                        ws.column_dimensions[letter].width = 44
                else:
                    ws.column_dimensions[letter].width = 28

    print(f"Planilha gerada: {path_out}")
    return 0

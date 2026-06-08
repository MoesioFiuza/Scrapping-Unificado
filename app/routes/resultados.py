from flask import Blueprint, request, jsonify, send_file, session
import pandas as pd
import os
from datetime import datetime
import traceback
from pathlib import Path
from app.services.transformador_dados import TransformadorDados
from app.services.extracoes_service import ExtracoesService
from app.utils.auth_decorator import login_required

bp = Blueprint('resultados', __name__, url_prefix='/api')

@bp.route('/resultados/exportar', methods=['POST'])
@login_required
def exportar_resultados():
    data = request.get_json()
    resultados = data.get('resultados', [])
    
    if not resultados:
        return jsonify({'error': 'Nenhum resultado para exportar'}), 400
    
    try:
        dados_export = []
        for resultado in resultados:
            if resultado.get('status') == 'sucesso' and resultado.get('dados'):
                dados = resultado['dados']
                dados_processo = dados.get('dados_processo', {})
                polo_ativo = dados.get('polo_ativo', [])
                polo_passivo = dados.get('polo_passivo', [])
                movimentacoes = dados.get('movimentacoes', [])
                texto_movimentacoes = '\n'.join([
                    mov.get('movimento', '') or mov.get('descricao', '')
                    for mov in movimentacoes
                    if mov.get('movimento') or mov.get('descricao')
                ])
                
                def limpar_texto(texto):
                    if not texto:
                        return ""
                    import re
                    texto = re.sub(r'\s+', ' ', str(texto))
                    return texto.strip()
                
                nomes_polo_ativo = [TransformadorDados.normalizar_nome_participante(p.get('nome', '') or p.get('nome_completo', '')) for p in polo_ativo]
                nomes_polo_passivo = [TransformadorDados.normalizar_nome_participante(p.get('nome', '') or p.get('nome_completo', '')) for p in polo_passivo]
                dados_export.append({
                    'Numero Processo': limpar_texto(resultado.get('numero_processo', '')),
                    'Tribunal': limpar_texto(resultado.get('tribunal', '')),
                    'Data Distribuicao': limpar_texto(dados_processo.get('data_distribuicao', '')),
                    'Classe Judicial': limpar_texto(dados_processo.get('classe_judicial', '')),
                    'Assunto': limpar_texto(dados_processo.get('assunto', '')),
                    'Jurisdicao': limpar_texto(dados_processo.get('jurisdicao', '')),
                    'Orgao Julgador': limpar_texto(dados_processo.get('orgao_julgador', '')),
                    'Polo Ativo': '; '.join(n for n in nomes_polo_ativo if n),
                    'Polo Passivo': '; '.join(n for n in nomes_polo_passivo if n),
                    'Movimentacoes': limpar_texto(texto_movimentacoes),
                    'Total Movimentacoes': len(movimentacoes),
                    'Total Documentos': len(dados.get('documentos', []))
                })
        
        if not dados_export:
            return jsonify({'error': 'Nenhum resultado com sucesso para exportar'}), 400
        df = pd.DataFrame(dados_export)
        
        username = session.get('username', '')
        downloads_path = Path.home() / 'Downloads'
        downloads_path.mkdir(exist_ok=True)
        
        output_dir = ExtracoesService.OUTPUT_DIR
        output_dir.mkdir(parents=True, exist_ok=True)
        
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f"resultados_raspados_{timestamp}.xlsx"
        
        downloads_filepath = downloads_path / filename
        with pd.ExcelWriter(str(downloads_filepath), engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Resultados')
            worksheet = writer.sheets['Resultados']
            worksheet.column_dimensions['A'].width = 20  # Numero Processo
            worksheet.column_dimensions['B'].width = 15  # Tribunal
            worksheet.column_dimensions['C'].width = 15  # Data Distribuicao
            worksheet.column_dimensions['D'].width = 40  # Classe Judicial
            worksheet.column_dimensions['E'].width = 50  # Assunto
            worksheet.column_dimensions['F'].width = 30  # Jurisdicao
            worksheet.column_dimensions['G'].width = 40  # Orgao Julgador
            worksheet.column_dimensions['H'].width = 40  # Polo Ativo
            worksheet.column_dimensions['I'].width = 40  # Polo Passivo
            worksheet.column_dimensions['J'].width = 80  # Movimentacoes (mais larga para o texto)
            worksheet.column_dimensions['K'].width = 15  # Total Movimentacoes
            worksheet.column_dimensions['L'].width = 15  # Total Documentos
            from openpyxl.styles import Alignment
            for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
        
        output_filepath = output_dir / filename
        with pd.ExcelWriter(str(output_filepath), engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Resultados')
            worksheet = writer.sheets['Resultados']
            worksheet.column_dimensions['A'].width = 20
            worksheet.column_dimensions['B'].width = 15
            worksheet.column_dimensions['C'].width = 15
            worksheet.column_dimensions['D'].width = 40
            worksheet.column_dimensions['E'].width = 50
            worksheet.column_dimensions['F'].width = 30
            worksheet.column_dimensions['G'].width = 40
            worksheet.column_dimensions['H'].width = 40
            worksheet.column_dimensions['I'].width = 40
            worksheet.column_dimensions['J'].width = 80
            worksheet.column_dimensions['K'].width = 15
            worksheet.column_dimensions['L'].width = 15
            from openpyxl.styles import Alignment
            for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
        
        total_processos = len(dados_export)
        ExtracoesService.registrar_extracao(
            username=username,
            filename=filename,
            tipo='raspado',
            total_processos=total_processos
        )
        
        return send_file(
            str(downloads_filepath), 
            as_attachment=True, 
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        
    except Exception as e:
        error_trace = traceback.format_exc()
        print(f"Erro ao exportar resultados: {error_trace}")
        return jsonify({'error': f'Erro ao exportar resultados: {str(e)}'}), 500

@bp.route('/resultados/exportar-tratado', methods=['POST'])
@login_required
def exportar_resultados_tratados():
    data = request.get_json()
    resultados = data.get('resultados', [])
    
    if not resultados:
        return jsonify({'error': 'Nenhum resultado para exportar'}), 400
    
    try:
        dados_transformados = TransformadorDados.transformar_lote(resultados)
        
        if not dados_transformados:
            return jsonify({'error': 'Nenhum resultado válido para transformar'}), 400
        
        df = pd.DataFrame(dados_transformados)
        
        # Coluna Advogado Parte Contraria (nome + OAB) entre partePoloAtivo e tipoPartePoloPassivo
        colunas_ordenadas = [
            'pasta', 'numeroProcessoAnterior', 'cnj', 'tipoPartePoloAtivo', 'partePoloAtivo',
            'Advogado Parte Contraria',
            'tipoPartePoloPassivo', 'partePoloPassivo', 'cliente', 'tipoDeRito', 'dataDistribuicao',
            'numeroUnidade', 'unidade', 'especialidade', 'comarca', 'estado', 'orgao', 'natureza',
            'materia', 'dataInstancia', 'tipoInstancia', 'sistemaExterno', 'processoEletronico',
            'processoEstrategico', 'valorCausa', 'valorFinalCausa', 'tipoAcao', 'tipoObjeto',
            'dataFase', 'fase', 'dataStatus', 'status', 'grupoProcesso', 'prioridadeDe',
            'data_resultado', 'tipo_resultado', 'descricao_resultado', 'dataEvento', 'tipoEvento',
            'descricaoEvento', 'complementoEvento', 'observacaoEvento', 'solicitanteEvento',
            'responsavelEvento', 'grupoTrabalho', 'corresponsavel', 'dataNotificacao',
            'dataNotificacaoAdicional', 'probabilidadePerda', 'dataValorProvisionado',
            'valorProvisionado', 'dataAndamento', 'tipoAndamento', 'descricaoAndamento',
            'complementoAndamento', 'solicitanteAndamento', 'responsavelAndamento',
            'corresponsavelAndamento', 'descricaoObjeto', 'escritorioCredenciado',
            'dataContratacao', 'observacaoDoProcesso', 'parecerDoProcesso'
        ]
        
        colunas_existentes = [col for col in colunas_ordenadas if col in df.columns]
        df = df[colunas_existentes]
        
        username = session.get('username', '')
        downloads_path = Path.home() / 'Downloads'
        downloads_path.mkdir(exist_ok=True)
        
        output_dir = ExtracoesService.OUTPUT_DIR
        output_dir.mkdir(parents=True, exist_ok=True)
        
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f"resultados_tratados_{timestamp}.xlsx"
        
        downloads_filepath = downloads_path / filename
        
        with pd.ExcelWriter(str(downloads_filepath), engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Processos')
            worksheet = writer.sheets['Processos']
            
            from openpyxl.utils import get_column_letter
            from openpyxl.styles import Alignment
            
            colunas_sem_quebra = ['cnj', 'numeroProcessoAnterior', 'tipoPartePoloAtivo', 'tipoPartePoloPassivo', 
                                 'tipoDeRito', 'dataDistribuicao', 'numeroUnidade', 'especialidade', 'comarca', 
                                 'estado', 'natureza', 'materia', 'tipoInstancia', 'processoEletronico', 
                                 'processoEstrategico', 'tipoAcao', 'dataStatus', 'status', 'tipoEvento']
            
            for idx, col in enumerate(df.columns, 1):
                col_letter = get_column_letter(idx)
                if col == 'cnj':
                    worksheet.column_dimensions[col_letter].width = 25  # Número do processo mais largo
                elif col in ['partePoloAtivo', 'partePoloPassivo', 'Advogado Parte Contraria', 'descricaoEvento', 'descricaoAndamento']:
                    worksheet.column_dimensions[col_letter].width = 40  # Colunas de texto mais largas
                else:
                    worksheet.column_dimensions[col_letter].width = 20
            
            for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
                for idx, cell in enumerate(row, 1):
                    col_name = df.columns[idx - 1] if idx <= len(df.columns) else ''
                    wrap_text = col_name not in colunas_sem_quebra
                    cell.alignment = Alignment(wrap_text=wrap_text, vertical='top')
        
        output_filepath = output_dir / filename
        with pd.ExcelWriter(str(output_filepath), engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Processos')
            worksheet = writer.sheets['Processos']
            
            from openpyxl.utils import get_column_letter
            from openpyxl.styles import Alignment
            
            colunas_sem_quebra = ['cnj', 'numeroProcessoAnterior', 'tipoPartePoloAtivo', 'tipoPartePoloPassivo', 
                                 'tipoDeRito', 'dataDistribuicao', 'numeroUnidade', 'especialidade', 'comarca', 
                                 'estado', 'natureza', 'materia', 'tipoInstancia', 'processoEletronico', 
                                 'processoEstrategico', 'tipoAcao', 'dataStatus', 'status', 'tipoEvento']
            
            for idx, col in enumerate(df.columns, 1):
                col_letter = get_column_letter(idx)
                if col == 'cnj':
                    worksheet.column_dimensions[col_letter].width = 25
                elif col in ['partePoloAtivo', 'partePoloPassivo', 'Advogado Parte Contraria', 'descricaoEvento', 'descricaoAndamento']:
                    worksheet.column_dimensions[col_letter].width = 40
                else:
                    worksheet.column_dimensions[col_letter].width = 20
            
            for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
                for idx, cell in enumerate(row, 1):
                    col_name = df.columns[idx - 1] if idx <= len(df.columns) else ''
                    wrap_text = col_name not in colunas_sem_quebra
                    cell.alignment = Alignment(wrap_text=wrap_text, vertical='top')
        
        total_processos = len(dados_transformados)
        ExtracoesService.registrar_extracao(
            username=username,
            filename=filename,
            tipo='tratado',
            total_processos=total_processos
        )
        
        return send_file(
            str(downloads_filepath), 
            as_attachment=True, 
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        
    except Exception as e:
        error_trace = traceback.format_exc()
        print(f"Erro ao exportar resultados tratados: {error_trace}")
        return jsonify({'error': f'Erro ao exportar resultados tratados: {str(e)}'}), 500